from fastapi import APIRouter, Depends, HTTPException, Query, Request, status, UploadFile, File
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional, List
import uuid
import shutil
import csv
import io
from pathlib import Path

from database import get_db
from modules.products.models import Product, Collection, CollectionProduct
from modules.products.catalogue_sync import sync_product_to_bot
from modules.whatsapp.models import CompanyKnowledge
from modules.products.schemas import (
    ProductCreate, ProductResponse, ProductUpdate, ProductListResponse,
    CollectionCreate, CollectionResponse, CollectionUpdate, CollectionListResponse,
    InventoryListResponse, InventoryItemResponse, InventoryUpdate
)
from utils.cache import cache, CacheConfig, generate_cache_key, cache_invalidator as invalidator
from utils.jwt_auth import get_current_user_optional
from modules.users.access import require_store_capability


async def _authorize_products_request(request: Request, db: Session = Depends(get_db),
                                      current_user=Depends(get_current_user_optional)):
    path = request.url.path
    if request.method == 'GET' and '/images/' in path:
        return
    if not current_user:
        raise HTTPException(status_code=401, detail='Not authenticated')
    params = request.path_params
    store_id = params.get('store_id') or request.query_params.get('store_id')
    resource_stores = []
    for key, model in (('product_id', Product), ('collection_id', Collection)):
        if params.get(key):
            try:
                entity = db.query(model).filter(model.id == uuid.UUID(str(params[key]))).first()
            except ValueError:
                entity = None
            if not entity:
                raise HTTPException(status_code=404, detail='Product resource not found')
            resource_stores.append(entity.store_id)
    if resource_stores:
        if any(value != resource_stores[0] for value in resource_stores):
            raise HTTPException(status_code=403, detail='Resources belong to different stores')
        store_id = resource_stores[0]
    if params.get('sku'):
        entity = db.query(Product).filter(Product.sku == params['sku']).first()
        if not entity:
            raise HTTPException(status_code=404, detail='Product not found')
        store_id = entity.store_id
    payload = None
    if request.method in ('POST', 'PUT', 'PATCH') and 'application/json' in request.headers.get('content-type', ''):
        try:
            payload = await request.json()
        except ValueError:
            pass
    if not store_id and isinstance(payload, dict):
        store_id = payload.get('store_id')
    # The product and collection list queries otherwise return every store.
    if request.method == 'GET' and path.rstrip('/').endswith(('/products', '/products/collections')) and not store_id:
        raise HTTPException(status_code=400, detail='Store is required')
    store_id = store_id or current_user.store_id
    if not store_id:
        raise HTTPException(status_code=403, detail='Store could not be determined')
    store = require_store_capability(db, current_user, store_id, 'products')
    # Reject mixed-store bulk updates and cross-store collection membership.
    rows = payload if isinstance(payload, list) else ([payload] if isinstance(payload, dict) else [])
    for row in rows:
        if not isinstance(row, dict):
            continue
        ids = list(row.get('product_ids') or []) + ([row['product_id']] if row.get('product_id') else [])
        for item_id in ids:
            try:
                item = db.query(Product).filter(Product.id == uuid.UUID(str(item_id))).first()
            except ValueError:
                item = None
            if not item or item.store_id != store.id:
                raise HTTPException(status_code=403, detail='Product belongs to another store')
        for collection_id in row.get('collections') or []:
            try:
                collection = db.query(Collection).filter(Collection.id == uuid.UUID(str(collection_id))).first()
            except ValueError:
                collection = None
            if not collection or collection.store_id != store.id:
                raise HTTPException(status_code=403, detail='Collection belongs to another store')


products_router = APIRouter(prefix="/api/products", tags=["products"],
                            dependencies=[Depends(_authorize_products_request)])

# Image upload configuration
UPLOAD_DIR = Path("uploads")
PRODUCT_IMAGE_DIR = UPLOAD_DIR / "products"
COLLECTION_IMAGE_DIR = UPLOAD_DIR / "collections"

# Create upload directories if they don't exist
PRODUCT_IMAGE_DIR.mkdir(parents=True, exist_ok=True)
COLLECTION_IMAGE_DIR.mkdir(parents=True, exist_ok=True)

# =============================================================================
# INVENTORY ENDPOINTS
# =============================================================================

# GET all inventory items (all products with inventory data)
@products_router.get("/inventory/store/{store_id}/items", response_model=InventoryListResponse)
async def get_inventory_items(
    store_id: uuid.UUID,
    skip: int = Query(0, ge=0),
    limit: int = Query(10, ge=1, le=100),
    product_id: Optional[uuid.UUID] = None,
    sku: Optional[str] = None,
    search: Optional[str] = None,
    status: Optional[str] = None,
    db: Session = Depends(get_db)
):
    try:
        # Build query with filters - store_id is REQUIRED for multi-tenant isolation
        query = db.query(Product).filter(Product.store_id == store_id)

        if product_id:
            query = query.filter(Product.id == product_id)
        if sku:
            query = query.filter(Product.sku.ilike(f"%{sku}%"))
        if status:
            query = query.filter(Product.status == status)
        if search:
            query = query.filter(
                Product.title.ilike(f"%{search}%") |
                Product.sku.ilike(f"%{search}%")
            )
        
        # Get total count
        total = query.count()
        
        # Apply pagination
        products = query.order_by(Product.created_at.asc(), Product.id.asc()).offset(skip).limit(limit).all()
        
        # Convert to inventory items format
        inventory_items = []
        for product in products:
            inventory_items.append(InventoryItemResponse(
                id=product.id,
                name=product.title,
                sku=product.sku,
                quantity=product.quantity,  # Add quantity
                unavailable=product.unavailable,
                committed=product.committed,
                available=product.available,
                on_hand=product.on_hand,
                status=product.status,
                inventory_tracked=product.inventory_tracked
            ))
        
        return InventoryListResponse(
            inventory_items=inventory_items,
            total=total,
            page=skip // limit + 1,
            size=len(inventory_items)
        )
        
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error fetching inventory items: {str(e)}"
        )

# GET single product inventory
@products_router.get("/inventory/items/{product_id}", response_model=InventoryItemResponse)
async def get_inventory_item(product_id: uuid.UUID, db: Session = Depends(get_db)):
    try:
        product = db.query(Product).filter(Product.id == product_id).first()
        
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
            
        return InventoryItemResponse(
            id=product.id,
            name=product.title,
            sku=product.sku,
            unavailable=product.unavailable,
            committed=product.committed,
            available=product.available,
            on_hand=product.on_hand,
            status=product.status,
            inventory_tracked=product.inventory_tracked
        )
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error fetching inventory item: {str(e)}"
        )

# UPDATE product inventory quantities
@products_router.put("/inventory/items/{product_id}", response_model=InventoryItemResponse)
async def update_inventory_item(
    product_id: uuid.UUID, 
    inventory_data: InventoryUpdate, 
    db: Session = Depends(get_db)
):
    try:
        # Get existing product
        product = db.query(Product).filter(Product.id == product_id).first()
        
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
        
        # Update inventory fields
        update_data = inventory_data.model_dump(exclude_unset=True)
        for field, value in update_data.items():
            setattr(product, field, value)
        
        db.commit()
        db.refresh(product)
        invalidator.invalidate_product()  # stock changed -> bust cached listings

        return InventoryItemResponse(
            id=product.id,
            name=product.title,
            sku=product.sku,
            quantity=product.quantity,  # Add quantity
            unavailable=product.unavailable,
            committed=product.committed,
            available=product.available,
            on_hand=product.on_hand,
            status=product.status,
            inventory_tracked=product.inventory_tracked
        )
        
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error updating inventory item: {str(e)}"
        )

# BULK UPDATE inventory quantities
@products_router.put("/inventory/bulk-update", response_model=List[InventoryItemResponse])
async def bulk_update_inventory(
    updates: List[dict],
    db: Session = Depends(get_db)
):
    try:
        results = []
        for update in updates:
            product_id = update.get('product_id')
            if not product_id:
                continue
                
            product = db.query(Product).filter(Product.id == product_id).first()
            if not product:
                continue
                
            # Update inventory fields
            if 'unavailable' in update:
                product.unavailable = update['unavailable']
            if 'committed' in update:
                product.committed = update['committed']
            if 'available' in update:
                product.available = update['available']
            if 'on_hand' in update:
                product.on_hand = update['on_hand']
            
            results.append(InventoryItemResponse(
                id=product.id,
                name=product.title,
                sku=product.sku,
                unavailable=product.unavailable,
                committed=product.committed,
                available=product.available,
                on_hand=product.on_hand,
                status=product.status,
                inventory_tracked=product.inventory_tracked
            ))
        
        db.commit()
        invalidator.invalidate_product()  # stock changed -> bust cached listings
        return results

    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error bulk updating inventory: {str(e)}"
        )

# UPDATE inventory quantities (quick update)
@products_router.patch("/inventory/items/{product_id}/quantities", response_model=InventoryItemResponse)
async def update_inventory_quantities(
    product_id: uuid.UUID,
    unavailable: Optional[int] = None,
    committed: Optional[int] = None,
    available: Optional[int] = None,
    on_hand: Optional[int] = None,
    db: Session = Depends(get_db)
):
    try:
        # Get existing product
        product = db.query(Product).filter(Product.id == product_id).first()
        
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
        
        # Update quantities if provided
        if unavailable is not None:
            product.unavailable = unavailable
        if committed is not None:
            product.committed = committed
        if available is not None:
            product.available = available
        if on_hand is not None:
            product.on_hand = on_hand
        
        db.commit()
        db.refresh(product)
        invalidator.invalidate_product()  # stock changed -> bust cached listings

        return InventoryItemResponse(
            id=product.id,
            name=product.title,
            sku=product.sku,
            unavailable=product.unavailable,
            committed=product.committed,
            available=product.available,
            on_hand=product.on_hand,
            status=product.status,
            inventory_tracked=product.inventory_tracked
        )

    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error updating inventory quantities: {str(e)}"
        )

# =============================================================================
# IMAGE UPLOAD ENDPOINTS
# =============================================================================

@products_router.post("/upload/product-image/{product_id}")
async def upload_product_image(
    product_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    try:
        # Verify product exists
        product = db.query(Product).filter(Product.id == product_id).first()
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
        
        # Validate file type
        allowed_extensions = {'.jpg', '.jpeg', '.png', '.gif', '.webp'}
        file_extension = Path(file.filename).suffix.lower()
        if file_extension not in allowed_extensions:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid file type. Allowed types: jpg, jpeg, png, gif, webp"
            )
        
        # Generate unique filename
        unique_filename = f"{product_id}_{uuid.uuid4()}{file_extension}"
        file_path = PRODUCT_IMAGE_DIR / unique_filename
        
        # Save file
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        
        # Update product with image paths
        product.image_path = str(file_path)
        product.image_url = f"/api/products/images/products/{unique_filename}"
        
        db.commit()
        db.refresh(product)
        
        return {
            "message": "Image uploaded successfully",
            "image_url": product.image_url,
            "image_path": product.image_path
        }
        
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error uploading image: {str(e)}"
        )

@products_router.post("/upload/collection-image/{collection_id}")
async def upload_collection_image(
    collection_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    try:
        # Verify collection exists
        collection = db.query(Collection).filter(Collection.id == collection_id).first()
        if not collection:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Collection not found"
            )
        
        # Validate file type
        allowed_extensions = {'.jpg', '.jpeg', '.png', '.gif', '.webp'}
        file_extension = Path(file.filename).suffix.lower()
        if file_extension not in allowed_extensions:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid file type. Allowed types: jpg, jpeg, png, gif, webp"
            )
        
        # Generate unique filename
        unique_filename = f"{collection_id}_{uuid.uuid4()}{file_extension}"
        file_path = COLLECTION_IMAGE_DIR / unique_filename
        
        # Save file
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        
        # Update collection with image paths
        collection.image_path = str(file_path)
        collection.image_url = f"/api/products/images/collections/{unique_filename}"
        
        db.commit()
        db.refresh(collection)
        
        return {
            "message": "Image uploaded successfully",
            "image_url": collection.image_url,
            "image_path": collection.image_path
        }
        
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error uploading image: {str(e)}"
        )

# Serve uploaded images
@products_router.get("/images/products/{filename}")
async def get_product_image(filename: str):
    file_path = PRODUCT_IMAGE_DIR / filename
    if not file_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Image not found"
        )
    return FileResponse(file_path)

@products_router.get("/images/collections/{filename}")
async def get_collection_image(filename: str):
    file_path = COLLECTION_IMAGE_DIR / filename
    if not file_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Image not found"
        )
    return FileResponse(file_path)

# =============================================================================
# PRODUCT ENDPOINTS
# =============================================================================

# GET all products with pagination and filtering
@products_router.get("/", response_model=ProductListResponse)
async def get_products(
    skip: int = Query(0, ge=0),
    limit: int = Query(10, ge=1, le=5000),
    store_id: Optional[uuid.UUID] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    vendor: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db)
):
    try:
        # Check cache
        cache_key = generate_cache_key(
            "products", skip, limit, store_id, category, status, vendor, search,
            prefix="prod:"
        )
        cached = cache.get(cache_key)
        if cached is not None:
            return cached

        # Build query with filters
        query = db.query(Product)

        # Filter by store_id (required for multi-tenant)
        if store_id:
            query = query.filter(Product.store_id == store_id)

        if category:
            query = query.filter(Product.category == category)
        if status:
            query = query.filter(Product.status == status)
        if vendor:
            query = query.filter(Product.vendor == vendor)
        if search:
            query = query.filter(
                Product.title.ilike(f"%{search}%") |
                Product.description.ilike(f"%{search}%") |
                Product.sku.ilike(f"%{search}%")
            )

        # Get total count
        total = query.count()

        # Apply pagination
        products = query.order_by(Product.created_at.asc(), Product.id.asc()).offset(skip).limit(limit).all()

        result = ProductListResponse(
            products=products,
            total=total,
            page=skip // limit + 1,
            size=len(products)
        )
        cache.set(cache_key, result.model_dump(), CacheConfig.TTL_MEDIUM)
        return result

    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error fetching products: {str(e)}"
        )

# =============================================================================
# PRODUCT IMPORT / EXPORT (must be before /{product_id} to avoid route conflict)
# =============================================================================

@products_router.get("/template/download")
async def download_product_template():
    """Download Excel template for product import"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    headers = [
        "title", "description", "category", "branch", "in_stock", "price", "sku", "barcode",
        "quantity", "weight", "weight_unit", "status", "vendor", "tags", "type", "image_url"
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = max(15, len(header) + 5)

    sample_data = [
        "Sample Product", "A great product description", "Clothing", "Main", "Yes", 29.99,
        "SKU-001", "123456789", 100, 0.5, "kg", "Active", "BrandName",
        "summer, new", "Physical", ""
    ]

    for col, value in enumerate(sample_data, 1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    instructions_ws = wb.create_sheet(title="Instructions")
    instructions = [
        ("Column", "Description", "Required", "Example"),
        ("title", "Product name", "Yes", "Blue T-Shirt"),
        ("description", "Product description", "No", "Cotton t-shirt"),
        ("category", "Product category", "No", "Clothing"),
        ("branch", "Branch name", "No", "Main"),
        ("in_stock", "Yes, No, or leave blank if unknown", "No", "Yes"),
        ("price", "Product price", "Yes", "29.99"),
        ("sku", "Stock keeping unit code", "No", "BTS-001"),
        ("barcode", "Product barcode", "No", "123456789"),
        ("quantity", "Stock quantity", "No", "100"),
        ("weight", "Product weight", "No", "0.5"),
        ("weight_unit", "Weight unit (kg/lb/g/oz)", "No", "kg"),
        ("status", "Product status (Active/Draft/Archived)", "No", "Active"),
        ("vendor", "Vendor or brand name", "No", "Nike"),
        ("tags", "Comma-separated tags", "No", "summer, sale"),
        ("type", "Product type", "No", "Physical"),
        ("image_url", "URL to product image", "No", "https://example.com/image.jpg"),
    ]

    for col, value in enumerate(instructions[0], 1):
        cell = instructions_ws.cell(row=1, column=col, value=value)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, row_data in enumerate(instructions[1:], 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = instructions_ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx == 3 and value == "Yes":
                cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    instructions_ws.column_dimensions['A'].width = 15
    instructions_ws.column_dimensions['B'].width = 40
    instructions_ws.column_dimensions['C'].width = 10
    instructions_ws.column_dimensions['D'].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=product_import_template.xlsx"}
    )


@products_router.post("/bulk-upload")
async def bulk_upload_products(
    file: UploadFile = File(...),
    store_id: uuid.UUID = Query(...),
    db: Session = Depends(get_db)
):
    """Bulk upload products from CSV or Excel file"""
    filename = file.filename.lower()

    if not (filename.endswith('.csv') or filename.endswith('.xlsx') or filename.endswith('.xls')):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only CSV and Excel (.xlsx, .xls) files are supported"
        )

    content = await file.read()
    created_products = []
    updated_products = []
    failed_products = []
    rows_to_process = []
    # A branch and SKU identify a row; products without a SKU use their title.
    # Ambiguous existing keys are never silently merged.
    existing_index = {}
    for existing in db.query(Product).filter(Product.store_id == store_id).all():
        branch_key = (existing.branch or '').strip().casefold()
        sku = (existing.sku or '').strip()
        identity = ('sku', sku.casefold()) if sku and sku.lower() != 'no sku' else ('title', existing.title.strip().casefold())
        key = (branch_key, *identity)
        existing_index[key] = None if key in existing_index else existing

    if filename.endswith('.csv'):
        csv_content = io.StringIO(content.decode('utf-8'))
        csv_reader = csv.DictReader(csv_content)
        for row_num, row in enumerate(csv_reader, start=2):
            rows_to_process.append((row_num, row))
    else:
        try:
            from openpyxl import load_workbook
            excel_content = io.BytesIO(content)
            workbook = load_workbook(excel_content, read_only=True)
            ws = workbook.active

            headers = []
            for cell in ws[1]:
                headers.append(cell.value.lower().strip() if cell.value else '')

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row):
                    continue
                row_dict = {}
                for col_idx, value in enumerate(row):
                    if col_idx < len(headers) and headers[col_idx]:
                        row_dict[headers[col_idx]] = str(value) if value is not None else None
                if row_dict.get('title') == 'Sample Product':
                    continue
                rows_to_process.append((row_num, row_dict))

            workbook.close()
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Error parsing Excel file: {str(e)}"
            )

    for row_num, row in rows_to_process:
        try:
            title = (row.get('title') or '').strip()
            price_str = (row.get('price') or '').strip()

            if not title or not price_str:
                failed_products.append({
                    "row": row_num,
                    "title": title or 'N/A',
                    "error": "Missing required fields (title, price)"
                })
                continue

            try:
                price = float(price_str)
            except ValueError:
                failed_products.append({
                    "row": row_num,
                    "title": title,
                    "error": f"Invalid price: {price_str}"
                })
                continue

            quantity = 0
            if row.get('quantity'):
                try:
                    quantity = int(float(row['quantity']))
                except (ValueError, TypeError):
                    pass

            if ((row.get('in_stock') or '').strip().lower() in ('yes', 'true')
                    and (row.get('quantity') or '').strip() and quantity <= 0):
                failed_products.append({'row': row_num, 'title': title,
                                        'error': 'In-stock products must have a positive quantity when quantity is provided'})
                continue
            stock_text = (row.get('in_stock') or '').strip().lower()
            if stock_text and stock_text not in ('yes', 'no', 'true', 'false'):
                failed_products.append({'row': row_num, 'title': title, 'error': 'in_stock must be Yes, No or blank'})
                continue

            weight = 0.0
            if row.get('weight'):
                try:
                    weight = float(row['weight'])
                except (ValueError, TypeError):
                    pass

            branch = (row.get('branch') or '').strip() or None
            sku = (row.get('sku') or '').strip() or 'No SKU'
            identity = ('sku', sku.casefold()) if sku.lower() != 'no sku' else ('title', title.casefold())
            key = ((branch or '').casefold(), *identity)
            if key in existing_index and existing_index[key] is None:
                failed_products.append({'row': row_num, 'title': title, 'error': 'Multiple existing products match this branch and SKU'})
                continue
            product = existing_index.get(key)
            if product:
                if (stock_text in ('yes', 'true') and product.inventory_tracked
                        and not (row.get('quantity') or '').strip()
                        and not (product.available or product.on_hand)):
                    failed_products.append({'row': row_num, 'title': title,
                        'error': 'Enter a positive quantity for this tracked product'})
                    continue
                product.title, product.price = title, price
                for field in ('description', 'category', 'barcode', 'vendor', 'tags', 'type', 'image_url', 'weight_unit', 'status'):
                    value = row.get(field)
                    if value is not None and str(value).strip():
                        setattr(product, field, str(value).strip())
                if stock_text:
                    product.in_stock = stock_text in ('yes', 'true')
                if (row.get('quantity') or '').strip():
                    product.quantity = product.available = product.on_hand = quantity
                    product.inventory_tracked = True
                if (row.get('weight') or '').strip():
                    product.weight = weight
                updated_products.append(product)
                continue
            product = Product(
                store_id=store_id,
                title=title,
                description=row.get('description'),
                category=row.get('category'),
                branch=branch,
                in_stock={'yes': True, 'true': True, 'no': False, 'false': False}.get(stock_text),
                price=price,
                sku=sku,
                barcode=row.get('barcode'),
                quantity=quantity,
                available=quantity,
                on_hand=quantity,
                inventory_tracked=bool((row.get('quantity') or '').strip()),
                weight=weight,
                weight_unit=row.get('weight_unit') or 'kg',
                status=row.get('status') or 'Active',
                vendor=row.get('vendor'),
                tags=row.get('tags'),
                type=row.get('type'),
                image_url=row.get('image_url'),
            )
            db.add(product)
            existing_index[key] = product
            created_products.append(product)

        except Exception as e:
            failed_products.append({
                "row": row_num,
                "title": row.get('title', 'N/A'),
                "error": str(e)
            })

    db.flush()
    unique_products = {product.id: product for product in created_products + updated_products}
    for product in unique_products.values():
        sync_product_to_bot(db, product)
    db.commit()
    invalidator.invalidate_product()
    for product in created_products:
        db.refresh(product)

    return {
        "created_count": len(created_products),
        "updated_count": len(updated_products),
        "failed_count": len(failed_products),
        "failed_products": failed_products,
    }

# =============================================================================
# PRODUCT DRAFTS
# =============================================================================

@products_router.get("/drafts/latest")
async def get_latest_draft(
    store_id: uuid.UUID = Query(...),
    db: Session = Depends(get_db)
):
    """Get the most recently updated draft product for a store"""
    draft = db.query(Product).filter(
        Product.store_id == store_id,
        Product.status == "Draft"
    ).order_by(Product.updated_at.desc()).first()

    if not draft:
        return None

    return draft

# =============================================================================
# PRODUCT CRUD
# =============================================================================

# GET single product by ID
@products_router.get("/{product_id}", response_model=ProductResponse)
async def get_product(product_id: uuid.UUID, db: Session = Depends(get_db)):
    try:
        product = db.query(Product).filter(Product.id == product_id).first()
        
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
            
        return product
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error fetching product: {str(e)}"
        )

# GET product by SKU
@products_router.get("/sku/{sku}", response_model=ProductResponse)
async def get_product_by_sku(sku: str, db: Session = Depends(get_db)):
    try:
        product = db.query(Product).filter(Product.sku == sku).first()
        
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
            
        return product
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error fetching product: {str(e)}"
        )

# CREATE product
# CREATE product
@products_router.post("/", response_model=ProductResponse, status_code=status.HTTP_201_CREATED)
async def create_product(product_data: ProductCreate, db: Session = Depends(get_db)):
    try:
        # Prepare product data - handle empty strings by converting to None
        product_dict = product_data.model_dump(exclude={'collections'})
        
        # Convert all empty strings to None for database compatibility
        for key, value in product_dict.items():
            if isinstance(value, str) and value.strip() == '':
                product_dict[key] = None
        
        # Special handling for SKU
        if not product_dict.get('sku'):
            product_dict['sku'] = 'No SKU'
        if product_dict.get('quantity') and not product_dict.get('available') and not product_dict.get('on_hand'):
            product_dict['available'] = product_dict['quantity']
            product_dict['on_hand'] = product_dict['quantity']
        if (product_dict.get('in_stock') is True and product_dict.get('inventory_tracked')
                and not (product_dict.get('available') or product_dict.get('on_hand'))):
            raise HTTPException(status_code=400, detail='Enter a positive quantity or disable inventory tracking for an in-stock item')
        
        # Check if SKU already exists (only if not using default)
        if product_dict['sku'] != 'No SKU':
            existing_product = db.query(Product).filter(
                Product.sku == product_dict['sku'], Product.store_id == product_dict['store_id'],
                Product.branch == product_dict.get('branch')
            ).first()
            
            if existing_product:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="SKU already exists"
                )
        
        # Create new product
        db_product = Product(**product_dict)
        db.add(db_product)
        db.flush()
        sync_product_to_bot(db, db_product)
        
        # Handle collections if provided
        if product_data.collections:
            for collection_id in product_data.collections:
                collection_product = CollectionProduct(
                    product_id=db_product.id,
                    collection_id=collection_id
                )
                db.add(collection_product)
        
        db.commit()
        db.refresh(db_product)
        invalidator.invalidate_product()

        return db_product

    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        print(f"Detailed error: {str(e)}")  # Add this for debugging
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error creating product: {str(e)}"
        )

# UPDATE product
@products_router.put("/{product_id}", response_model=ProductResponse)
async def update_product(
    product_id: uuid.UUID, 
    product_data: ProductUpdate, 
    db: Session = Depends(get_db)
):
    try:
        # Get existing product
        db_product = db.query(Product).filter(Product.id == product_id).first()
        
        if not db_product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
        
        # Prepare update data - handle empty strings
        update_data = product_data.model_dump(exclude_unset=True)
        
        # Convert empty strings to None for nullable fields
        nullable_fields = ['description', 'category', 'type', 'vendor', 'tags', 'barcode', 'image_url', 'image_path']
        for field in nullable_fields:
            if field in update_data and update_data[field] == '':
                update_data[field] = None
        
        # Handle SKU specifically
        if 'sku' in update_data and (not update_data['sku'] or update_data['sku'].strip() == ''):
            update_data['sku'] = 'No SKU'
        
        # Check if SKU already exists (if being updated and not using default)
        if 'sku' in update_data and update_data['sku'] != 'No SKU' and update_data['sku'] != db_product.sku:
            existing_product = db.query(Product).filter(
                Product.sku == update_data['sku'], Product.store_id == db_product.store_id,
                Product.branch == update_data.get('branch', db_product.branch), Product.id != db_product.id
            ).first()
            
            if existing_product:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="SKU already exists"
                )
        
        # Update fields
        for field, value in update_data.items():
            setattr(db_product, field, value)
        if 'quantity' in update_data and 'available' not in update_data and 'on_hand' not in update_data:
            db_product.available = update_data['quantity']
            db_product.on_hand = update_data['quantity']
        if (db_product.in_stock is True and db_product.inventory_tracked
                and not (db_product.available or db_product.on_hand)):
            raise HTTPException(status_code=400, detail='Enter a positive quantity or disable inventory tracking for an in-stock item')
        sync_product_to_bot(db, db_product)
        
        db.commit()
        db.refresh(db_product)
        invalidator.invalidate_product()

        return db_product

    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error updating product: {str(e)}"
        )

# DELETE product
@products_router.delete("/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_product(product_id: uuid.UUID, db: Session = Depends(get_db)):
    try:
        # Get existing product
        db_product = db.query(Product).filter(Product.id == product_id).first()
        
        if not db_product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
        
        # Delete associated collection relationships first
        db.query(CollectionProduct).filter(
            CollectionProduct.product_id == product_id
        ).delete()
        db.query(CompanyKnowledge).filter(CompanyKnowledge.store_id == db_product.store_id,
            CompanyKnowledge.source_id == db_product.id, CompanyKnowledge.content_type == 'product').update(
                {'is_active': False}, synchronize_session=False)
        
        db.delete(db_product)
        db.commit()
        invalidator.invalidate_product()

        return None

    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error deleting product: {str(e)}"
        )

# =============================================================================
# COLLECTION ENDPOINTS
# =============================================================================

# GET all collections with pagination and filtering
@products_router.get("/collections/", response_model=CollectionListResponse)
async def get_collections(
    skip: int = Query(0, ge=0),
    limit: int = Query(10, ge=1, le=100),
    store_id: Optional[uuid.UUID] = None,
    collection_type: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db)
):
    try:
        # Check cache
        cache_key = generate_cache_key(
            "collections", skip, limit, store_id, collection_type, search,
            prefix="prod:"
        )
        cached = cache.get(cache_key)
        if cached is not None:
            return cached

        # Build query with filters
        query = db.query(Collection)

        # Filter by store_id (required for multi-tenant)
        if store_id:
            query = query.filter(Collection.store_id == store_id)

        if collection_type:
            query = query.filter(Collection.collection_type == collection_type)
        if search:
            query = query.filter(Collection.title.ilike(f"%{search}%"))

        # Get total count
        total = query.count()

        # Apply pagination
        collections = query.offset(skip).limit(limit).all()

        result = CollectionListResponse(
            collections=collections,
            total=total,
            page=skip // limit + 1,
            size=len(collections)
        )
        cache.set(cache_key, result.model_dump(), CacheConfig.TTL_MEDIUM)
        return result

    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error fetching collections: {str(e)}"
        )

# GET single collection by ID
@products_router.get("/collections/{collection_id}", response_model=CollectionResponse)
async def get_collection(collection_id: uuid.UUID, db: Session = Depends(get_db)):
    try:
        collection = db.query(Collection).filter(Collection.id == collection_id).first()
        
        if not collection:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Collection not found"
            )
            
        return collection
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error fetching collection: {str(e)}"
        )

# CREATE collection
@products_router.post("/collections/", response_model=CollectionResponse, status_code=status.HTTP_201_CREATED)
async def create_collection(collection_data: CollectionCreate, db: Session = Depends(get_db)):
    try:
        # Check if collection title already exists
        existing_collection = db.query(Collection).filter(
            Collection.title == collection_data.title,
            Collection.store_id == collection_data.store_id
        ).first()
        
        if existing_collection:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Collection title already exists"
            )
        
        # Create new collection
        collection_dict = collection_data.model_dump(exclude={'product_ids'})
        db_collection = Collection(**collection_dict)
        db.add(db_collection)
        db.flush()  # Flush to get the collection ID
        
        # Handle products if provided
        if collection_data.product_ids:
            for product_id in collection_data.product_ids:
                # Verify product exists
                product = db.query(Product).filter(Product.id == product_id).first()
                if not product:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"Product with ID {product_id} not found"
                    )
                
                collection_product = CollectionProduct(
                    collection_id=db_collection.id,
                    product_id=product_id
                )
                db.add(collection_product)
        
        db.commit()
        db.refresh(db_collection)
        invalidator.invalidate_product()

        return db_collection

    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error creating collection: {str(e)}"
        )

# UPDATE collection
@products_router.put("/collections/{collection_id}", response_model=CollectionResponse)
async def update_collection(
    collection_id: uuid.UUID, 
    collection_data: CollectionUpdate, 
    db: Session = Depends(get_db)
):
    try:
        # Get existing collection
        db_collection = db.query(Collection).filter(Collection.id == collection_id).first()
        
        if not db_collection:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Collection not found"
            )
        
        # Check if title already exists (if being updated)
        if collection_data.title and collection_data.title != db_collection.title:
            existing_collection = db.query(Collection).filter(
                Collection.title == collection_data.title
            ).first()
            
            if existing_collection:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Collection title already exists"
                )
        
        # Update fields
        update_data = collection_data.model_dump(exclude_unset=True)
        for field, value in update_data.items():
            setattr(db_collection, field, value)
        
        db.commit()
        db.refresh(db_collection)
        invalidator.invalidate_product()

        return db_collection

    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error updating collection: {str(e)}"
        )

# DELETE collection
@products_router.delete("/collections/{collection_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_collection(collection_id: uuid.UUID, db: Session = Depends(get_db)):
    try:
        # Get existing collection
        db_collection = db.query(Collection).filter(Collection.id == collection_id).first()
        
        if not db_collection:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Collection not found"
            )
        
        # Delete associated collection relationships first
        db.query(CollectionProduct).filter(
            CollectionProduct.collection_id == collection_id
        ).delete()
        
        db.delete(db_collection)
        db.commit()
        invalidator.invalidate_product()

        return None

    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error deleting collection: {str(e)}"
        )

# ADD product to collection
@products_router.post("/collections/{collection_id}/products/{product_id}", response_model=CollectionResponse)
async def add_product_to_collection(
    collection_id: uuid.UUID,
    product_id: uuid.UUID,
    db: Session = Depends(get_db)
):
    try:
        # Verify collection exists
        collection = db.query(Collection).filter(Collection.id == collection_id).first()
        if not collection:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Collection not found"
            )
        
        # Verify product exists
        product = db.query(Product).filter(Product.id == product_id).first()
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found"
            )
        
        # Check if product already in collection
        existing_relation = db.query(CollectionProduct).filter(
            CollectionProduct.collection_id == collection_id,
            CollectionProduct.product_id == product_id
        ).first()
        
        if existing_relation:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Product already in collection"
            )
        
        # Add product to collection
        collection_product = CollectionProduct(
            collection_id=collection_id,
            product_id=product_id
        )
        db.add(collection_product)
        db.commit()
        db.refresh(collection)
        
        return collection
        
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error adding product to collection: {str(e)}"
        )

# REMOVE product from collection
@products_router.delete("/collections/{collection_id}/products/{product_id}", response_model=CollectionResponse)
async def remove_product_from_collection(
    collection_id: uuid.UUID,
    product_id: uuid.UUID,
    db: Session = Depends(get_db)
):
    try:
        # Verify collection exists
        collection = db.query(Collection).filter(Collection.id == collection_id).first()
        if not collection:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Collection not found"
            )
        
        # Remove product from collection
        db.query(CollectionProduct).filter(
            CollectionProduct.collection_id == collection_id,
            CollectionProduct.product_id == product_id
        ).delete()
        
        db.commit()
        db.refresh(collection)
        
        return collection
        
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error removing product from collection: {str(e)}"
        )

# GET collection products
@products_router.get("/collections/{collection_id}/products", response_model=ProductListResponse)
async def get_collection_products(
    collection_id: uuid.UUID,
    skip: int = Query(0, ge=0),
    limit: int = Query(10, ge=1, le=100),
    db: Session = Depends(get_db)
):
    try:
        # Verify collection exists
        collection = db.query(Collection).filter(Collection.id == collection_id).first()
        if not collection:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Collection not found"
            )
        
        # Get products in collection
        query = db.query(Product).join(
            CollectionProduct, 
            Product.id == CollectionProduct.product_id
        ).filter(CollectionProduct.collection_id == collection_id)
        
        total = query.count()
        products = query.offset(skip).limit(limit).all()
        
        return ProductListResponse(
            products=products,
            total=total,
            page=skip // limit + 1,
            size=len(products)
        )
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error fetching collection products: {str(e)}"
        )




